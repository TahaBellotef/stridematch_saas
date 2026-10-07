#
#  File: services/dashboard_export.py
#  Project: StrideMatchLab
#
#  Excel (.xlsx) and PDF export builders for the admin dashboard tabs.
#

from __future__ import annotations

import io
from typing import TYPE_CHECKING

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

if TYPE_CHECKING:
    from app.schemas.dashboard import (
        DashboardCustomersResponse,
        DashboardInsightsResponse,
        DashboardOverviewResponse,
    )

_HEADER_FONT = Font(bold=True)


def _write_table(ws: Worksheet, headers: list, rows: list) -> None:
    ws.append(headers)
    for cell in ws[ws.max_row]:
        cell.font = _HEADER_FONT
    for row in rows:
        ws.append(row)
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(40, width + 2)


def _workbook_to_bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_customers_workbook(data: "DashboardCustomersResponse") -> bytes:
    wb = Workbook()

    summary = wb.active
    summary.title = "Summary"
    _write_table(
        summary,
        ["Metric", "Value", "Change vs previous period"],
        [
            ["New Customers", data.new_customers.display, data.new_customers.change_pct],
            ["Returning Customers", data.returning_customers.display, data.returning_customers.change_pct],
            ["Retention Rate", data.retention_rate.display, data.retention_rate.change_pct],
            ["Total Customers", data.total_customers, None],
        ],
    )

    trend = wb.create_sheet("Acquisition vs Retention")
    _write_table(
        trend,
        ["Period", "Date", "Acquisition", "Retention"],
        [[p.label, p.date, p.acquisition, p.retention] for p in data.daily_acquisition_retention],
    )

    experience = wb.create_sheet("Experience")
    _write_table(
        experience,
        ["Experience", "Customers"],
        [[seg.label, seg.count] for seg in data.experience_breakdown],
    )

    frequency = wb.create_sheet("Run Frequency")
    _write_table(
        frequency,
        ["Run Frequency", "Customers"],
        [[seg.label, seg.count] for seg in data.run_frequency_breakdown],
    )

    return _workbook_to_bytes(wb)


def build_insights_workbook(data: "DashboardInsightsResponse") -> bytes:
    wb = Workbook()

    summary = wb.active
    summary.title = "Summary"
    _write_table(
        summary,
        ["Metric", "Value", "Change vs previous period"],
        [
            ["Total Analyses", data.total_analyses.display, data.total_analyses.change_pct],
            ["Back View", data.back_view_count.display, data.back_view_count.change_pct],
            ["Side View", data.side_view_count.display, data.side_view_count.change_pct],
            ["Total Scans", data.total_scans, None],
            ["Total Recommendations", data.total_recommendations, None],
        ],
    )

    analysis_breakdown = wb.create_sheet("Analysis by Period")
    _write_table(
        analysis_breakdown,
        ["Period", "Date", "Back View", "Side View", "Unspecified"],
        [[p.label, p.date, p.rear, p.side, p.unspecified] for p in data.daily_analysis_breakdown],
    )

    scans = wb.create_sheet("Scans by Period")
    _write_table(
        scans,
        ["Period", "Date", "Scans"],
        [[p.label, p.date, p.value] for p in data.daily_scan_counts],
    )

    recommendations = wb.create_sheet("Recommendations by Period")
    _write_table(
        recommendations,
        ["Period", "Date", "Recommendations"],
        [[p.label, p.date, p.value] for p in data.daily_recommendation_counts],
    )

    type_breakdown = wb.create_sheet("Analysis Type Breakdown")
    _write_table(
        type_breakdown,
        ["Type", "Count"],
        [[seg.label, seg.count] for seg in data.analysis_type_breakdown],
    )

    return _workbook_to_bytes(wb)


def build_overview_pdf(data: "DashboardOverviewResponse", period_label: str) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter, title="StrideMatch Dashboard Overview")
    styles = getSampleStyleSheet()
    story = [
        Paragraph("StrideMatch - Dashboard Overview", styles["Title"]),
        Paragraph(f"Period: {period_label}", styles["Normal"]),
        Spacer(1, 16),
    ]

    def _change_text(change_pct):
        if change_pct is None:
            return "-"
        sign = "+" if change_pct >= 0 else ""
        return f"{sign}{change_pct}%"

    metrics_table = Table(
        [
            ["Metric", "Value", "Change vs previous period"],
            ["Scans (this period)", data.scans_this_week.display, _change_text(data.scans_this_week.change_pct)],
            ["Total Scans (all-time)", data.total_scans.display, "-"],
            [
                "New Customers (this period)",
                data.new_customers_this_week.display,
                _change_text(data.new_customers_this_week.change_pct),
            ],
            ["Conversion Rate", data.conversion_rate.display, _change_text(data.conversion_rate.change_pct)],
            ["Avg. Confidence", f"{data.avg_confidence_pct}%" if data.avg_confidence_pct is not None else "-", "-"],
            ["Total Customers", str(data.total_customers), "-"],
        ],
        hAlign="LEFT",
    )
    metrics_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4B21EF")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F5F9")]),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    story.append(metrics_table)
    story.append(Spacer(1, 20))

    story.append(Paragraph("Customer Demographics", styles["Heading2"]))
    demographics_table = Table(
        [["Segment", "Count"]] + [[seg.segment.title(), str(seg.count)] for seg in data.demographics],
        hAlign="LEFT",
    )
    demographics_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#312D4B")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    story.append(demographics_table)
    story.append(Spacer(1, 20))

    if data.activity:
        story.append(Paragraph("Recent Activity", styles["Heading2"]))
        activity_rows = [["Title", "Customer", "When"]] + [
            [item.title, item.customer_name or "-", item.timestamp] for item in data.activity
        ]
        activity_table = Table(activity_rows, hAlign="LEFT")
        activity_table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#312D4B")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                    ("TOPPADDING", (0, 0), (-1, -1), 5),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                ]
            )
        )
        story.append(activity_table)

    doc.build(story)
    return buf.getvalue()
